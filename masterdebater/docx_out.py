"""Write the generated speech to a Word document."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path


def _add_markdownish(doc, text: str) -> None:
    """Very small Markdown subset: #/##/### headings, '- ' bullets, **bold** runs."""
    from docx.shared import Pt

    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            doc.add_heading(m.group(2).strip(), level=min(level, 3))
            continue
        if re.match(r"^\s*[-*]\s+", line):
            para = doc.add_paragraph(style="List Bullet")
            body = re.sub(r"^\s*[-*]\s+", "", line)
        else:
            para = doc.add_paragraph()
            body = line
        for chunk in re.split(r"(\*\*[^*]+\*\*)", body):
            if not chunk:
                continue
            if chunk.startswith("**") and chunk.endswith("**"):
                run = para.add_run(chunk[2:-2])
                run.bold = True
            else:
                run = para.add_run(chunk)
            run.font.size = Pt(12)


def write_speech_docx(
    out_path: str | Path,
    speech_name: str,
    side_label: str,
    speech_text: str,
    notes: str | None = None,
    word_count: int | None = None,
    minutes: float | None = None,
) -> Path:
    from docx import Document

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading(f"{speech_name} — {side_label}", level=0)
    meta = [datetime.now().strftime("%Y-%m-%d %H:%M")]
    if minutes:
        meta.append(f"{minutes:g} min")
    if word_count:
        meta.append(f"{word_count} words")
    doc.add_paragraph(" · ".join(meta)).italic = True

    _add_markdownish(doc, speech_text)

    if notes:
        doc.add_page_break()
        doc.add_heading("Coach notes (not part of the speech)", level=1)
        _add_markdownish(doc, notes)

    doc.save(str(out_path))
    return out_path
