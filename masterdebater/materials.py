"""Load the debater's own case, blocks and evidence files.

These files are the ONLY source material the speech writer may quote from. We accept
Word (.docx), PDF, plain text and Markdown. Each file is returned as a labelled block so
the model can cite which file a card came from.
"""

from __future__ import annotations

from pathlib import Path

SUPPORTED = {".docx", ".pdf", ".txt", ".md", ".markdown"}


def _read_docx(path: Path) -> str:
    from docx import Document  # python-docx

    doc = Document(str(path))
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def read_material(path: str | Path) -> str:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".docx":
        return _read_docx(path)
    if ext == ".pdf":
        return _read_pdf(path)
    if ext in SUPPORTED:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported material type {ext!r}; use one of {sorted(SUPPORTED)}")


def read_material_bytes(name: str, data: bytes) -> str:
    """Same as read_material but for an in-memory upload (Streamlit)."""
    import io
    import tempfile

    ext = Path(name).suffix.lower()
    if ext in {".txt", ".md", ".markdown"}:
        return data.decode("utf-8", errors="replace")
    if ext == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    if ext == ".pdf":
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp.write(data)
            tmp_path = tmp.name
        try:
            return _read_pdf(Path(tmp_path))
        finally:
            Path(tmp_path).unlink(missing_ok=True)
    raise ValueError(f"Unsupported material type {ext!r}; use one of {sorted(SUPPORTED)}")


def bundle_materials(named_texts: list[tuple[str, str]]) -> str:
    """Join several (filename, text) pairs into one labelled block for the prompt."""
    if not named_texts:
        return "(no case or block files were provided)"
    out = []
    for name, text in named_texts:
        out.append(f"<<< FILE: {name} >>>\n{text.strip()}\n<<< END FILE: {name} >>>")
    return "\n\n".join(out)
