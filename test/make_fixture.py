"""Generate a realistic debate-blocks .docx for tests (mirrors the user's AB1..ABn layout)."""
import sys
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor

out = sys.argv[1] if len(sys.argv) > 1 else "test/fixtures/blocks.docx"
doc = Document()
doc.add_heading("Neg Blocks — National Service", level=0)  # Title style, not a heading -> ignored preamble

def card(doc, tag, cite, text, highlighted_words=()):
    p = doc.add_paragraph()
    r = p.add_run(tag)
    r.bold = True
    r.font.size = Pt(12)
    p2 = doc.add_paragraph()
    r = p2.add_run(cite)
    r.underline = True
    r.font.size = Pt(10)
    p3 = doc.add_paragraph()
    for word in text.split(" "):
        r = p3.add_run(word + " ")
        r.font.size = Pt(8)
        if word.strip(",.") in highlighted_words:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            r.underline = True
            r.font.size = Pt(11)

blocks = [
    ("AB1 — Coercion / Instrumentalization", "Mandatory service treats people as means.", ["treats", "means"]),
    ("AB2 — 13th Amendment", "The draft was upheld in Arver v. United States in 1918.", ["upheld", "1918"]),
    ("AB3 — Coerced Virtue / Voluntary Is Better", "Virtue that is compelled is not virtue at all.", ["compelled"]),
    ("AB4 — Cost / Feasibility", "A universal program would cost over 100 billion dollars per year.", ["100", "billion"]),
    ("AB8 — Liberty / Mill's harm principle (PRIORITY)", "The only purpose for which power can be rightfully exercised is to prevent harm to others.", ["prevent", "harm"]),
]
for i, (title, cite_text, hl) in enumerate(blocks, start=1):
    doc.add_heading(title, level=1)
    p = doc.add_paragraph()
    r = p.add_run("CARDS: ")
    r.bold = True
    p.add_run("read the first two, cut the rest if short on time.")
    note = doc.add_paragraph()
    nr = note.add_run("NOTE: ")
    nr.bold = True
    nr.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    note.add_run("signpost before reading.")
    doc.add_heading("ONLY IF they go for the framework", level=2)
    card(doc, f"One. {title.split(' — ')[1]} outweighs.", "Author 2019, Journal of Things", cite_text, hl)
    if i == 2:
        t = doc.add_table(rows=2, cols=2)
        t.style = "Table Grid"
        t.cell(0, 0).text = "Year"
        t.cell(0, 1).text = "Case"
        t.cell(1, 0).text = "1918"
        t.cell(1, 1).text = "Arver"
    if i == 4:
        doc.add_paragraph("")  # empty paragraph inside a block

# A duplicate heading title, to exercise id de-duplication
doc.add_heading("AB4 — Cost / Feasibility", level=1)
doc.add_paragraph("Second copy of the cost block.")
# A heading with no body (should still be a block, since it is at the split level)
doc.add_heading("AB20 — Mental health harms", level=1)
doc.save(out)
print("wrote", out)
